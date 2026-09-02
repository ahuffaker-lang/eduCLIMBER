"""Reading source spreadsheets (CSV or Excel) and matching their headers loosely."""

from __future__ import annotations

import csv
import io
import os
import re
from dataclasses import dataclass
from typing import Dict, Iterable, List, Optional, Sequence

CSV_SUFFIXES = {".csv", ".txt", ".tsv"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


class SourceFileError(Exception):
    """The source file could not be read at all."""


def normalize_header(header: str) -> str:
    """Fold a header down to comparable letters/digits.

    ``"  Reading Level (A-Z) "`` and ``"reading_level_az"`` both become
    ``"readinglevelaz"``, which is what makes header matching best-effort rather
    than exact.
    """
    return re.sub(r"[^a-z0-9]+", "", (header or "").lower())


@dataclass
class Table:
    """A source file loaded into memory."""

    headers: List[str]
    rows: List[Dict[str, str]]
    first_data_row_number: int = 2  # 1-based; row 1 is the header

    def row_number(self, index: int) -> int:
        return self.first_data_row_number + index

    def find_column(self, candidates: Sequence[str]) -> Optional[str]:
        """Return the actual header matching any of *candidates*, else ``None``."""
        wanted = [normalize_header(c) for c in candidates]
        by_norm = {normalize_header(h): h for h in self.headers}
        for name in wanted:
            if name in by_norm:
                return by_norm[name]
        # Fall back to a containment match so a header like "Student ID Number"
        # still matches "Student ID". Short names ("id", "grade") are skipped --
        # they turn up inside unrelated headers and match the wrong column.
        for name in wanted:
            if len(name) < 8:
                continue
            for norm, original in by_norm.items():
                if name in norm or norm in name:
                    return original
        return None

    def value(self, row: Dict[str, str], column: Optional[str]) -> str:
        if not column:
            return ""
        raw = row.get(column, "")
        return "" if raw is None else str(raw).strip()


def read_table(path: str) -> Table:
    """Load a CSV or Excel file into a :class:`Table`."""
    suffix = os.path.splitext(path)[1].lower()
    if suffix in EXCEL_SUFFIXES:
        return _read_excel(path)
    if suffix in CSV_SUFFIXES or not suffix:
        return _read_csv(path)
    if suffix == ".xls":
        raise SourceFileError(
            "Legacy .xls files are not supported. Please re-save the export as "
            ".xlsx or .csv and try again."
        )
    raise SourceFileError(f"Unsupported file type '{suffix}'. Use a .csv or .xlsx file.")


def _read_csv(path: str) -> Table:
    try:
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(handle, dialect)
            rows = [row for row in reader]
    except UnicodeDecodeError:
        with open(path, "r", encoding="latin-1", newline="") as handle:
            rows = list(csv.reader(handle))
    except OSError as exc:
        raise SourceFileError(f"Could not open the file: {exc}") from exc
    return _from_grid(rows)


def _read_excel(path: str) -> Table:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise SourceFileError(
            "Reading Excel files needs the 'openpyxl' package. Either install it "
            "(pip install openpyxl) or save the export as a .csv file."
        ) from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a variety of types
        raise SourceFileError(f"Could not open the Excel file: {exc}") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        grid = [
            ["" if cell is None else _clean_cell(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()
    return _from_grid(grid)


def _clean_cell(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _from_grid(grid: Iterable[Sequence[object]]) -> Table:
    grid = [list(row) for row in grid]
    header_index = _find_header_row(grid)
    if header_index is None:
        raise SourceFileError("The file appears to be empty -- no header row found.")

    headers = _dedupe([str(cell).strip() for cell in grid[header_index]])
    rows: List[Dict[str, str]] = []
    for raw in grid[header_index + 1:]:
        values = ["" if cell is None else str(cell).strip() for cell in raw]
        if not any(values):
            continue  # skip blank spacer rows
        values += [""] * (len(headers) - len(values))
        rows.append(dict(zip(headers, values)))
    return Table(headers=headers, rows=rows, first_data_row_number=header_index + 2)


def _find_header_row(grid: Sequence[Sequence[object]]) -> Optional[int]:
    """The first non-blank row is the header; exports sometimes start with padding."""
    for index, row in enumerate(grid):
        if any(str(cell).strip() for cell in row if cell is not None):
            return index
    return None


def _dedupe(headers: Sequence[str]) -> List[str]:
    seen: Dict[str, int] = {}
    result: List[str] = []
    for position, header in enumerate(headers):
        name = header or f"Column {position + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        result.append(name)
    return result


def write_csv(path: str, columns: Sequence[str], rows: Iterable[Dict[str, object]]) -> None:
    """Write *rows* to *path* with exactly *columns*, in order.

    Plain UTF-8 with no byte-order mark: a BOM would land inside the first header
    name ("Assessment Group") when eduCLIMBER's importer parses the file.
    """
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})


def rows_to_csv_text(columns: Sequence[str], rows: Iterable[Dict[str, object]]) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(columns), extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({column: row.get(column, "") for column in columns})
    return buffer.getvalue()
